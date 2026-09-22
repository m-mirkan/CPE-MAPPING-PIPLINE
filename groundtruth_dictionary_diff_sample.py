from __future__ import annotations
import json
import logging
import os
from dataclasses import dataclass
from typing import List, Set, Tuple

from cpe_utils import parse_cpe23
from local_dictionary import LocalCpeDictionary

logger = logging.getLogger("cpe_pipeline.groundtruth_dictionary_diff_sample")


@dataclass
class Discrepancy:
    cpe: str
    role: str            # "hardware" or "firmware_os"
    source_cves: List[str]


def _cpe_covered_by_dictionary(cpe_string: str, local_dict: LocalCpeDictionary) -> bool:
   
    parsed = parse_cpe23(cpe_string)
    if parsed is None:
        return False
    if parsed.version == "*":
        return len(local_dict.entries_for(parsed.part, parsed.vendor, parsed.product)) > 0
    return local_dict.exact_exists(cpe_string) is not None


def device_coverage_ceiling_for_sample(
    groundtruth_raw: list, sample_hw_cpes: List[str], local_dict: LocalCpeDictionary,
) -> dict:
   
    sample_set = set(sample_hw_cpes)
    gt_by_hw = {e["hardware_cpe"]: e for e in groundtruth_raw if e.get("hardware_cpe")}

    with_groundtruth = [hw for hw in sample_set if hw in gt_by_hw]
    coverable, uncoverable = [], []

    for hw in with_groundtruth:
        entry = gt_by_hw[hw]
        os_cpes = entry.get("os_firmware_cpes", []) or []
        has_any_coverage = any(_cpe_covered_by_dictionary(cpe, local_dict) for cpe in os_cpes)
        (coverable if has_any_coverage else uncoverable).append(hw)

    return {
        "devices_with_groundtruth": len(with_groundtruth),
        "devices_dictionary_coverable": len(coverable),
        "devices_dictionary_uncoverable": len(uncoverable),
        "uncoverable_hw_cpes": sorted(uncoverable),
    }


def find_discrepancies_for_sample(
    groundtruth_raw: list, sample_hw_cpes: List[str], local_dict: LocalCpeDictionary,
) -> Tuple[List[Discrepancy], List[str], int]:
   
    sample_set = set(sample_hw_cpes)
    gt_by_hw = {e["hardware_cpe"]: e for e in groundtruth_raw if e.get("hardware_cpe")}

    seen: dict = {}  # cpe -> [role, set(cves)]
    sample_hw_with_no_groundtruth: Set[str] = set()

    for hw_cpe in sample_set:
        entry = gt_by_hw.get(hw_cpe)
        if entry is None:
            sample_hw_with_no_groundtruth.add(hw_cpe)
            continue

        cves = entry.get("source_cves", []) or []
        bucket = seen.setdefault(hw_cpe, ["hardware", set()])
        bucket[1].update(cves)
        for os_cpe in entry.get("os_firmware_cpes", []) or []:
            bucket = seen.setdefault(os_cpe, ["firmware_os", set()])
            bucket[1].update(cves)

    discrepancies = []
    for cpe, (role, cve_set) in seen.items():
        if not _cpe_covered_by_dictionary(cpe, local_dict):
            discrepancies.append(Discrepancy(cpe=cpe, role=role, source_cves=sorted(cve_set)))
    discrepancies.sort(key=lambda d: (d.role, d.cpe))

    sample_hw_with_groundtruth_count = len(sample_set) - len(sample_hw_with_no_groundtruth)
    return discrepancies, sorted(sample_hw_with_no_groundtruth), sample_hw_with_groundtruth_count


def write_discrepancies_json(discrepancies: List[Discrepancy], output_path: str,
                              sample_size: int, sample_hw_with_no_groundtruth: List[str],
                              sample_hw_with_groundtruth_count: int,
                              device_coverage_ceiling: dict = None) -> None:
    payload = {
        "summary": {
            "sample_size": sample_size,
            "sample_hw_with_groundtruth": sample_hw_with_groundtruth_count,
            "sample_hw_with_no_groundtruth_entry": len(sample_hw_with_no_groundtruth),
            "total_discrepancies": len(discrepancies),
            "hardware_discrepancies": sum(1 for d in discrepancies if d.role == "hardware"),
            "firmware_os_discrepancies": sum(1 for d in discrepancies if d.role == "firmware_os"),
            "note": (
                "STRING-level question: across all sample_hw_with_groundtruth devices, how many "
                "DISTINCT ground-truth CPE strings (hardware or firmware/OS) are missing from the "
                "local dictionary? A single device can have several ground-truth firmware entries, "
                "so this count is not directly comparable to a per-device figure -- see "
                "device_coverage_ceiling below for the per-device version of this question. A "
                "ground-truth CPE with version '*' is treated as covered if ANY concrete version "
                "of that vendor+product exists in the dictionary."
            ),
        },
        "sample_hw_cpes_with_no_groundtruth_entry": sample_hw_with_no_groundtruth,
        "discrepancies": [
            {"cpe": d.cpe, "role": d.role, "source_cves": d.source_cves}
            for d in discrepancies
        ],
    }

    if device_coverage_ceiling is not None:
        payload["device_coverage_ceiling"] = {
            **device_coverage_ceiling,
            "note": (
                "DEVICE-level question, computed independently of the discrepancies above and "
                "WITHOUT reading matched_cpes.json or our pipeline's own predictions at all: of "
                "the devices_with_groundtruth devices, how many have AT LEAST ONE ground-truth "
                "firmware/OS CPE present in the dictionary? This is a theoretical CEILING -- the "
                "maximum number of devices ANY matcher could possibly succeed on against this "
                "dictionary snapshot, since a correct match requires the true answer to be present "
                "in the dictionary to begin with. It does not measure whether OUR matcher actually "
                "found it -- that is a separate, already-reported question (test_results.json)."
            ),
        }

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(payload, f, indent=2, ensure_ascii=False)
    logger.info("Wrote %d sample-scoped ground-truth/dictionary discrepancies to %s",
                len(discrepancies), output_path)


def write_discrepancies_xlsx(discrepancies: List[Discrepancy], output_path: str,
                              sample_hw_with_no_groundtruth: List[str],
                              device_coverage_ceiling: dict = None) -> bool:
   
    try:
        import openpyxl
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.warning(
            "openpyxl is not installed, skipping .xlsx output (the .json output was "
            "still written with the same data). Install it with: pip install openpyxl"
        )
        return False

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Discrepancies"
    headers = ["CPE", "Role", "Source CVEs", "CVE Count"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for d in discrepancies:
        ws.append([d.cpe, d.role, ", ".join(d.source_cves), len(d.source_cves)])
    for i, width in enumerate([70, 14, 60, 12], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("No Ground Truth")
    ws2.append(["Hardware CPE (in sample, no ground-truth entry at all)"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for hw in sample_hw_with_no_groundtruth:
        ws2.append([hw])
    ws2.column_dimensions["A"].width = 70

    if device_coverage_ceiling is not None:
        ws3 = wb.create_sheet("Dictionary-Uncoverable")
        ws3.append(["Hardware CPE (has ground truth, but NO ground-truth firmware/OS CPE exists in the dictionary at all)"])
        for cell in ws3[1]:
            cell.font = Font(bold=True)
        for hw in device_coverage_ceiling.get("uncoverable_hw_cpes", []):
            ws3.append([hw])
        ws3.column_dimensions["A"].width = 90

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    wb.save(output_path)
    logger.info("Wrote %d sample-scoped ground-truth/dictionary discrepancies to %s",
                len(discrepancies), output_path)
    return True